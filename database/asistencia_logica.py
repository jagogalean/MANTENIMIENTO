"""
Lógica de asistencia y horas extras (sin Streamlit ni base de datos: se puede probar sola).

Reglas (las mismas de la planilla anterior, más el criterio de «hora real»):
  - Horas netas = (salida - ingreso, cruzando la medianoche si hace falta) - descanso + horas extra adicionales.
  - Las primeras horas del turno, hasta la jornada base (8 hs por defecto), son ordinarias.
    Lo que pasa de la jornada son horas extra: diurnas (turno diurno o normal) o nocturnas (turno nocturno).
  - Las horas trabajadas en domingo o feriado se pagan completas con el factor de domingo/feriado
    (diurno o nocturno según el turno) y no se cuentan como ordinarias ni como extras comunes.
  - Criterio de hora real: un turno que cruza la medianoche se reparte entre los dos días y cada tramo
    se trata según el día en que cae (sábado de noche hasta el domingo a la mañana = parte hábil + parte domingo).
    El descanso se reparte en proporción a las horas de cada tramo y las horas adicionales (que se
    cargan al extender el turno) caen en el último tramo. Con el criterio desactivado, todo el turno
    se trata según el día en que empieza (como la planilla anterior).
"""
import io
import math
from datetime import date, datetime, time, timedelta

CATEGORIAS = ["he_diurna", "he_nocturna", "domingo_diurno", "domingo_nocturno", "feriado_diurno", "feriado_nocturno"]
NOMBRE_CATEGORIA = {
    "he_diurna": "HE diurnas", "he_nocturna": "HE nocturnas",
    "domingo_diurno": "Domingo diurno", "domingo_nocturno": "Domingo nocturno",
    "feriado_diurno": "Feriado diurno", "feriado_nocturno": "Feriado nocturno",
}
FACTOR_CATEGORIA = {
    "he_diurna": "factor_he_diurna", "he_nocturna": "factor_he_nocturna",
    "domingo_diurno": "factor_domingo_diurno", "domingo_nocturno": "factor_domingo_nocturno",
    "feriado_diurno": "factor_feriado_diurno", "feriado_nocturno": "factor_feriado_nocturno",
}
PARAMS_DEFECTO = {
    "dias_base": 30.0,                   # días del mes para calcular el valor hora
    "horas_jornada_diurna": 8.0,
    "horas_jornada_nocturna": 8.0,
    "factor_he_diurna": 1.5,             # +50%
    "factor_he_nocturna": 2.3,           # +130%
    "factor_domingo_diurno": 2.0,        # +100%
    "factor_domingo_nocturno": 2.3,      # +130% (100% de domingo + 30% nocturno)
    "factor_feriado_diurno": 2.0,
    "factor_feriado_nocturno": 2.3,
    "tope_he_diarias": 3.0,              # alerta: más horas extra por día que esto
    "domingo_todo_el_turno": 1.0,        # 1 = todo el turno de domingo/feriado se paga con su factor
    "cruce_por_hora_real": 1.0,          # 1 = turnos que cruzan la medianoche se reparten por día
}
ESTADOS_MODALIDAD = ["Trabajo", "Libre", "Ausente", "Vacaciones/Reposo"]
TIPOS_RECARGO = ["DIURNA", "NOCTURNA", "NINGUNO"]
MODALIDADES_DEFECTO = [
    {"nombre": "Diurno", "ingreso": "06:00", "salida": "18:00", "descanso": 1.0, "tipo": "DIURNA", "estado": "Trabajo"},
    {"nombre": "Nocturno", "ingreso": "18:00", "salida": "06:00", "descanso": 1.0, "tipo": "NOCTURNA", "estado": "Trabajo"},
    {"nombre": "Normal", "ingreso": "08:00", "salida": "17:00", "descanso": 1.0, "tipo": "DIURNA", "estado": "Trabajo"},
    {"nombre": "Libre", "ingreso": None, "salida": None, "descanso": 0.0, "tipo": "NINGUNO", "estado": "Libre"},
    {"nombre": "Ausente", "ingreso": None, "salida": None, "descanso": 0.0, "tipo": "NINGUNO", "estado": "Ausente"},
    {"nombre": "Vacaciones / Reposo", "ingreso": None, "salida": None, "descanso": 0.0, "tipo": "NINGUNO", "estado": "Vacaciones/Reposo"},
]


# ----------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------
def _n(v, defecto=0.0):
    try:
        return float(v) if v not in (None, "") else defecto
    except (TypeError, ValueError):
        return defecto


def a_fecha(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.strptime(str(v)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def a_hora(v):
    """time, datetime, 'HH:MM' o 'HH:MM:SS' -> time (o None)."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, time):
        return v
    try:
        partes = [int(x) for x in str(v).split(":")[:3]]
        return time(*partes)
    except (ValueError, TypeError):
        return None


def hora_decimal(t):
    return t.hour + t.minute / 60 + t.second / 3600


def hora_texto(t):
    t = a_hora(t)
    return t.strftime("%H:%M") if t else ""


def redondear(x):
    """Redondeo de Excel (la mitad hacia arriba), no el del banquero de Python."""
    return int(math.floor(x + 0.5)) if x >= 0 else -int(math.floor(-x + 0.5))


def parametros_completos(params):
    return {**PARAMS_DEFECTO, **{k: _n(v) for k, v in (params or {}).items() if k in PARAMS_DEFECTO}}


def tipo_de_dia(fecha, feriados):
    """'FERIADO' (aunque caiga domingo), 'DOMINGO' o 'HABIL'."""
    if fecha in feriados:
        return "FERIADO"
    return "DOMINGO" if fecha.weekday() == 6 else "HABIL"


def valor_hora(empleado, params):
    P = parametros_completos(params)
    total = _n(empleado.get("salario_base")) + _n(empleado.get("adicional_fijo"))
    return total / P["dias_base"] / P["horas_jornada_diurna"] if P["dias_base"] and P["horas_jornada_diurna"] else 0.0


# ----------------------------------------------------------------
# Cálculo de un día
# ----------------------------------------------------------------
def calcular_dia(fecha, modalidad, ingreso, salida, descanso, he_d_adic, he_n_adic, params, feriados):
    """
    Devuelve:
      netas, ordinarias, horas {categoria: hs}, segmentos [{fecha, tipo_dia, horas_netas}],
      tipo_dia (el del día en que empieza), cruza (bool), avisos [texto]
    """
    P = parametros_completos(params)
    fecha = a_fecha(fecha)
    res = {"netas": 0.0, "ordinarias": 0.0, "horas": {c: 0.0 for c in CATEGORIAS}, "segmentos": [], "cruza": False,
           "tipo_dia": tipo_de_dia(fecha, feriados) if fecha else "HABIL", "avisos": []}
    if not fecha or not modalidad or modalidad.get("estado", "Trabajo") != "Trabajo":
        return res
    t_in, t_out = a_hora(ingreso), a_hora(salida)
    if t_in is None or t_out is None:
        res["avisos"].append("Falta la hora de ingreso o de salida.")
        return res
    nocturno = modalidad.get("tipo") == "NOCTURNA"
    h_in, h_out = hora_decimal(t_in), hora_decimal(t_out)
    duracion = (h_out - h_in) % 24
    if duracion == 0:
        res["avisos"].append("El ingreso y la salida son iguales: no se contaron horas.")
    descanso = _n(descanso)
    neto_base = max(0.0, round(duracion - descanso, 2)) if duracion else 0.0
    he_d_adic, he_n_adic = max(0.0, _n(he_d_adic)), max(0.0, _n(he_n_adic))
    res["netas"] = round(neto_base + he_d_adic + he_n_adic, 2)

    # Tramos: uno solo, o dos si el turno cruza la medianoche y se calcula por hora real
    tramos = [(fecha, duracion)]
    if P["cruce_por_hora_real"] and h_in + duracion > 24 + 1e-9:
        primero = 24 - h_in
        tramos = [(fecha, primero), (fecha + timedelta(days=1), duracion - primero)]
    res["cruza"] = len(tramos) > 1
    netos, acumulado = [], 0.0
    for i, (_, bruto) in enumerate(tramos):
        neto = round(neto_base * bruto / duracion, 2) if duracion else 0.0
        if i == len(tramos) - 1:
            neto = round(neto_base - acumulado, 2)          # el último tramo absorbe el redondeo
        netos.append(max(neto, 0.0))
        acumulado += neto

    jornada = P["horas_jornada_nocturna"] if nocturno else P["horas_jornada_diurna"]
    sufijo = "nocturno" if nocturno else "diurno"
    corrido = 0.0
    for (dia, _), neto in zip(tramos, netos):
        tipo = tipo_de_dia(dia, feriados)
        res["segmentos"].append({"fecha": dia, "tipo_dia": tipo, "horas_netas": neto})
        if tipo != "HABIL" and P["domingo_todo_el_turno"]:
            res["horas"][f"{tipo.lower()}_{sufijo}"] += neto
        else:
            ordinarias = min(neto, max(0.0, jornada - corrido))
            res["ordinarias"] += ordinarias
            res["horas"]["he_nocturna" if nocturno else "he_diurna"] += neto - ordinarias
        corrido += neto
    # Las horas adicionales se agregan al terminar el turno: caen en el último tramo
    ultimo = res["segmentos"][-1]["tipo_dia"]
    if ultimo != "HABIL" and P["domingo_todo_el_turno"]:
        res["horas"][f"{ultimo.lower()}_{sufijo}"] += he_d_adic + he_n_adic
    else:
        res["horas"]["he_diurna"] += he_d_adic
        res["horas"]["he_nocturna"] += he_n_adic
    res["ordinarias"] = round(res["ordinarias"], 2)
    res["horas"] = {c: round(v, 2) for c, v in res["horas"].items()}
    he_comunes = res["horas"]["he_diurna"] + res["horas"]["he_nocturna"]
    if he_comunes > P["tope_he_diarias"] + 1e-9:
        res["avisos"].append(f"{he_comunes:g} horas extra en el día (el tope es {P['tope_he_diarias']:g}).")
    return res


def modalidad_por_nombre(modalidades):
    return {m["nombre"]: m for m in modalidades}


def calcular_registro(registro, modalidades_por_nombre, params, feriados):
    """Aplica calcular_dia a un registro de la base (con ingreso/salida como texto o time)."""
    mod = modalidades_por_nombre.get(registro.get("modalidad"))
    return calcular_dia(registro.get("fecha"), mod, registro.get("ingreso"), registro.get("salida"), registro.get("descanso"),
                        registro.get("he_diurna_adic"), registro.get("he_nocturna_adic"), params, feriados)


# ----------------------------------------------------------------
# Liquidación por período
# ----------------------------------------------------------------
def liquidar(empleados, registros, modalidades, params, feriados, desde, hasta):
    """
    Una fila por empleado con sus horas y montos entre 'desde' y 'hasta' (ambos incluidos).
    Devuelve {"filas": [...], "totales": {...}, "avisos": [...], "detalle": {empleado_id: [días]}}.
    """
    P = parametros_completos(params)
    desde, hasta = a_fecha(desde), a_fecha(hasta)
    mods = modalidad_por_nombre(modalidades)
    por_empleado = {}
    for r in registros:
        d = a_fecha(r.get("fecha"))
        if d and desde <= d <= hasta:
            por_empleado.setdefault(r["empleado_id"], []).append(r)
    dias_periodo = (hasta - desde).days + 1
    filas, avisos, detalle = [], [], {}
    for e in sorted(empleados, key=lambda x: (str(x.get("area") or "").casefold(), str(x.get("nombre") or "").casefold())):
        regs = sorted(por_empleado.get(e["id"], []), key=lambda r: a_fecha(r["fecha"]))
        if not regs and e.get("activo") is False:
            continue
        vh = valor_hora(e, P)
        horas = {c: 0.0 for c in CATEGORIAS}
        dias_trab = ausencias = libres = vacaciones = 0
        dias_con_aviso, det = [], []
        for r in regs:
            calc = calcular_registro(r, mods, P, feriados)
            mod = mods.get(r.get("modalidad"), {})
            if calc["netas"] > 0:
                dias_trab += 1
            estado = mod.get("estado")
            ausencias += estado == "Ausente"
            libres += estado == "Libre"
            vacaciones += estado == "Vacaciones/Reposo"
            for c in CATEGORIAS:
                horas[c] += calc["horas"][c]
            if calc["avisos"]:
                dias_con_aviso.append((a_fecha(r["fecha"]), calc["avisos"]))
            det.append({"fecha": a_fecha(r["fecha"]), "modalidad": r.get("modalidad"), "calculo": calc})
        fila = {"empleado_id": e["id"], "empleado": e.get("nombre"), "area": e.get("area") or "", "salario": _n(e.get("salario_base")) + _n(e.get("adicional_fijo")),
                "valor_hora": vh, "dias_trabajados": dias_trab, "ausencias": int(ausencias), "libres": int(libres), "vacaciones": int(vacaciones),
                "dias_sin_registro": max(0, dias_periodo - len({a_fecha(r["fecha"]) for r in regs}))}
        total = 0
        for c in CATEGORIAS:
            cantidad = round(horas[c], 2)
            tarifa = vh * P[FACTOR_CATEGORIA[c]]
            monto = redondear(cantidad * tarifa)
            fila[f"{c}_horas"], fila[f"{c}_tarifa"], fila[f"{c}_monto"] = cantidad, tarifa, monto
            total += monto
        fila["total"] = total
        filas.append(fila)
        detalle[e["id"]] = det
        if fila["salario"] <= 0 and sum(horas.values()) > 0:
            avisos.append({"nivel": "alta", "texto": f"{e.get('nombre')} tiene {sum(horas.values()):g} horas extra pero no tiene salario cargado: no se pueden valorizar."})
        for d, textos in dias_con_aviso:
            avisos.append({"nivel": "media", "texto": f"{e.get('nombre')} · {d:%d/%m/%Y}: {' '.join(textos)}"})
        if regs and fila["dias_sin_registro"]:
            avisos.append({"nivel": "info", "texto": f"{e.get('nombre')}: {fila['dias_sin_registro']} día(s) del período sin ningún registro cargado."})
    totales = {"salario": sum(f["salario"] for f in filas), "dias_trabajados": sum(f["dias_trabajados"] for f in filas),
               "ausencias": sum(f["ausencias"] for f in filas), "total": sum(f["total"] for f in filas)}
    for c in CATEGORIAS:
        totales[f"{c}_horas"] = round(sum(f[f"{c}_horas"] for f in filas), 2)
        totales[f"{c}_monto"] = sum(f[f"{c}_monto"] for f in filas)
    orden = {"alta": 0, "media": 1, "info": 2}
    return {"filas": filas, "totales": totales, "avisos": sorted(avisos, key=lambda a: orden[a["nivel"]]), "detalle": detalle}


# ----------------------------------------------------------------
# Cargas rápidas (reemplazan las macros de la planilla)
# ----------------------------------------------------------------
def fechas_de_semana(lunes, dias=6):
    """Lunes a sábado a partir del lunes indicado."""
    return [lunes + timedelta(days=i) for i in range(dias)]


def registro_por_defecto(empleado_id, fecha, modalidad, semana=None):
    """Registro con los horarios de la modalidad (como hacían las macros)."""
    es_trabajo = modalidad.get("estado") == "Trabajo"
    nota = {"Diurno": "1 hora de descanso", "Nocturno": "1 hora de descanso", "Normal": "Horario normal"}.get(modalidad.get("nombre"), "")
    return {"empleado_id": empleado_id, "fecha": fecha.isoformat(), "modalidad": modalidad["nombre"],
            "ingreso": hora_texto(modalidad.get("ingreso")) if es_trabajo else None, "salida": hora_texto(modalidad.get("salida")) if es_trabajo else None,
            "descanso": _n(modalidad.get("descanso")) if es_trabajo else 0.0, "he_diurna_adic": 0.0, "he_nocturna_adic": 0.0,
            "semana": semana or None, "observaciones": nota or None}


# ----------------------------------------------------------------
# Lectura de la planilla Excel anterior (.xlsm / .xlsx)
# ----------------------------------------------------------------
HOJA_PARAMS, HOJA_REGISTRO = "1_PARAMETROS_Y_NOMINA", "2_REGISTRO_ASISTENCIA"
CLAVE_PARAM_PLANILLA = {6: "dias_base", 7: "horas_jornada_diurna", 8: "horas_jornada_nocturna", 9: "factor_he_diurna",
                        10: "factor_he_nocturna", 11: "factor_domingo_diurno", 12: "tope_he_diarias"}


def leer_planilla(contenido: bytes):
    """
    Lee la planilla de control de asistencia. Devuelve {empleados, modalidades, feriados, registros, parametros, avisos}
    o {"error": texto}. No toca nada: solo interpreta el archivo.
    """
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=False)
    except Exception as e:
        return {"error": f"No se pudo abrir el archivo: {e}"}
    if HOJA_PARAMS not in wb.sheetnames or HOJA_REGISTRO not in wb.sheetnames:
        return {"error": f"No encuentro las hojas «{HOJA_PARAMS}» y «{HOJA_REGISTRO}». ¿Es la planilla de control de asistencia?"}
    p, r = wb[HOJA_PARAMS], wb[HOJA_REGISTRO]
    avisos = []

    parametros = {}
    for fila, clave in CLAVE_PARAM_PLANILLA.items():
        v = p.cell(row=fila, column=2).value
        if isinstance(v, (int, float)):
            parametros[clave] = float(v)
    parametros["domingo_todo_el_turno"] = 1.0 if str(p.cell(row=13, column=2).value or "").strip().upper() == "SI" else 0.0
    if "factor_domingo_diurno" in parametros:                      # la planilla tiene un solo factor para domingo y feriado
        parametros["factor_feriado_diurno"] = parametros["factor_domingo_diurno"]

    empleados = []
    for fila in range(6, p.max_row + 1):
        if p.cell(row=fila, column=5).value in (None, ""):          # la nómina termina donde se acaban los números de ID
            break
        nombre = p.cell(row=fila, column=6).value
        if nombre and str(nombre).strip():
            empleados.append({"nombre": str(nombre).strip(), "salario_base": _n(p.cell(row=fila, column=7).value),
                              "adicional_fijo": _n(p.cell(row=fila, column=8).value)})

    modalidades = []
    for fila in range(17, 40):
        nombre = p.cell(row=fila, column=1).value
        if not nombre or not str(nombre).strip():
            continue
        nombre = str(nombre).strip()
        recargo = str(p.cell(row=fila, column=6).value or "NINGUNO").strip().upper()
        estado = {"Libre": "Libre", "Ausente": "Ausente"}.get(nombre, "Vacaciones/Reposo" if nombre.lower().startswith("vacac") else "Trabajo")
        if recargo == "NINGUNO" and estado == "Trabajo":
            estado = "Libre"
        modalidades.append({"nombre": nombre, "ingreso": hora_texto(p.cell(row=fila, column=2).value) or None,
                            "salida": hora_texto(p.cell(row=fila, column=3).value) or None, "descanso": _n(p.cell(row=fila, column=4).value),
                            "tipo": recargo if recargo in TIPOS_RECARGO else "NINGUNO", "estado": estado})

    feriados = []
    for fila in range(17, 60):
        f = a_fecha(p.cell(row=fila, column=8).value)
        if f:
            feriados.append({"fecha": f.isoformat(), "motivo": str(p.cell(row=fila, column=9).value or "").strip() or None})

    registros = []
    for fila in range(5, r.max_row + 1):
        fecha, emp = a_fecha(r.cell(row=fila, column=3).value), r.cell(row=fila, column=6).value
        if not fecha or not emp or not str(emp).strip():
            continue
        registros.append({"empleado": str(emp).strip(), "fecha": fecha.isoformat(), "semana": str(r.cell(row=fila, column=2).value or "").strip() or None,
                          "modalidad": str(r.cell(row=fila, column=7).value or "").strip() or None,
                          "ingreso": hora_texto(r.cell(row=fila, column=8).value) or None, "salida": hora_texto(r.cell(row=fila, column=9).value) or None,
                          "descanso": _n(r.cell(row=fila, column=10).value), "he_diurna_adic": _n(r.cell(row=fila, column=11).value),
                          "he_nocturna_adic": _n(r.cell(row=fila, column=12).value),
                          "observaciones": str(r.cell(row=fila, column=18).value or "").strip() or None})
    nombres_nomina = {e["nombre"] for e in empleados}
    for nombre in sorted({x["empleado"] for x in registros} - nombres_nomina):
        empleados.append({"nombre": nombre, "salario_base": 0.0, "adicional_fijo": 0.0, "sin_salario": True})
        avisos.append(f"«{nombre}» tiene registros pero no figura en la nómina de la planilla: se importa sin salario (cargalo después) y sus horas no se podrán valorizar hasta entonces.")
    mod_nombres = {m["nombre"] for m in modalidades}
    for nombre in sorted({x["modalidad"] for x in registros if x["modalidad"]} - mod_nombres):
        avisos.append(f"La modalidad «{nombre}» se usa en el registro pero no está en el catálogo de la planilla.")
    return {"empleados": empleados, "modalidades": modalidades, "feriados": feriados, "registros": registros, "parametros": parametros,
            "avisos": avisos}
