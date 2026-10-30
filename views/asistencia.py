import io
from datetime import date, time, timedelta

import pandas as pd
import streamlit as st

from database import asistencia_db as db
from database import asistencia_logica as A

DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
ICONO = {"alta": "🔴", "media": "🟠", "info": "ℹ️"}
NOMBRE_PARAM = {
    "dias_base": "Días base del mes (para el valor hora)", "horas_jornada_diurna": "Horas de la jornada diurna",
    "horas_jornada_nocturna": "Horas de la jornada nocturna", "factor_he_diurna": "Hora extra diurna (factor)",
    "factor_he_nocturna": "Hora extra nocturna (factor)", "factor_domingo_diurno": "Domingo diurno (factor)",
    "factor_domingo_nocturno": "Domingo nocturno (factor)", "factor_feriado_diurno": "Feriado diurno (factor)",
    "factor_feriado_nocturno": "Feriado nocturno (factor)", "tope_he_diarias": "Alerta: más de estas horas extra por día",
}
CAMPOS_REGISTRO = ["modalidad", "ingreso", "salida", "descanso", "he_diurna_adic", "he_nocturna_adic", "observaciones"]


# ----------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------
def _fmt_gs(v):
    return "Gs. " + f"{v:,.0f}".replace(",", ".")


def _hs(v):
    return f"{v:g}".replace(".", ",")


TIPO_DIA_TXT = {"HABIL": "hábil", "DOMINGO": "domingo", "FERIADO": "feriado"}


def _tramos_txt(calculo):
    """Ej: «24/10 hábil 5,5 h + 25/10 domingo 5,5 h»"""
    return " + ".join(f"{s['fecha']:%d/%m} {TIPO_DIA_TXT[s['tipo_dia']]} {_hs(s['horas_netas'])} h" for s in calculo["segmentos"])


def _fecha_larga(d):
    return f"{DIAS[d.weekday()]} {d:%d/%m/%Y}"


def _hora_celda(v):
    """Valor de una celda de hora de la tabla editable -> time o None."""
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return A.a_hora(v)


def _normalizar(campo, v):
    if campo in ("ingreso", "salida"):
        t = _hora_celda(v)
        return t.strftime("%H:%M") if t else None
    if campo in ("descanso", "he_diurna_adic", "he_nocturna_adic"):
        try:
            return round(float(v), 2) if v is not None and not pd.isna(v) else 0.0
        except (TypeError, ValueError):
            return 0.0
    texto = "" if v is None else ("" if (isinstance(v, float) and pd.isna(v)) else str(v).strip())
    return texto or None


def calcular_cambios_registros(originales: dict, editados: dict):
    """{id: {campo: valor}} antes y después -> [(id, parche)] solo con lo que cambió (horas como 'HH:MM')."""
    cambios = []
    for rid, nuevo in editados.items():
        previo = originales.get(rid, {})
        parche = {}
        for campo in CAMPOS_REGISTRO:
            antes = _normalizar(campo, previo.get(campo))
            ahora = _normalizar(campo, nuevo.get(campo))
            if antes != ahora:
                parche[campo] = ahora
        if parche:
            cambios.append((rid, parche))
    return cambios


def construir_excel(liq, empleados_por_id, params, desde, hasta):
    """Excel con la liquidación, el detalle día por día y los parámetros usados (para archivar el cierre)."""
    filas = []
    for f in liq["filas"]:
        fila = {"Empleado": f["empleado"], "Área": f["area"], "Salario (Gs.)": f["salario"], "Valor hora (Gs.)": round(f["valor_hora"], 2),
                "Días trabajados": f["dias_trabajados"], "Ausencias": f["ausencias"], "Libres": f["libres"], "Vacaciones/Reposo": f["vacaciones"]}
        for c in A.CATEGORIAS:
            fila[f"{A.NOMBRE_CATEGORIA[c]} (hs)"] = f[f"{c}_horas"]
            fila[f"{A.NOMBRE_CATEGORIA[c]} (Gs.)"] = f[f"{c}_monto"]
        fila["TOTAL A PAGAR (Gs.)"] = f["total"]
        filas.append(fila)
    t = liq["totales"]
    total = {"Empleado": "TOTAL DEL PERÍODO", "Salario (Gs.)": t["salario"], "Días trabajados": t["dias_trabajados"], "Ausencias": t["ausencias"]}
    for c in A.CATEGORIAS:
        total[f"{A.NOMBRE_CATEGORIA[c]} (hs)"] = t[f"{c}_horas"]
        total[f"{A.NOMBRE_CATEGORIA[c]} (Gs.)"] = t[f"{c}_monto"]
    total["TOTAL A PAGAR (Gs.)"] = t["total"]
    filas.append(total)

    detalle = []
    for eid, dias in liq["detalle"].items():
        for d in dias:
            c = d["calculo"]
            if not c["netas"] and not d["modalidad"]:
                continue
            detalle.append({"Empleado": empleados_por_id[eid]["nombre"], "Fecha": d["fecha"], "Día": DIAS[d["fecha"].weekday()], "Modalidad": d["modalidad"],
                            "Tramos": _tramos_txt(c),
                            "Horas netas": c["netas"], "Ordinarias": c["ordinarias"],
                            **{A.NOMBRE_CATEGORIA[k]: v for k, v in c["horas"].items()}, "Avisos": " ".join(c["avisos"])})
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as w:
        pd.DataFrame(filas).to_excel(w, index=False, sheet_name="Liquidación")
        pd.DataFrame(detalle).to_excel(w, index=False, sheet_name="Detalle por día")
        pd.DataFrame([{"Parámetro": NOMBRE_PARAM.get(k, k), "Valor": v} for k, v in A.parametros_completos(params).items()]
                     + [{"Parámetro": "Período liquidado", "Valor": f"{desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"}]).to_excel(w, index=False, sheet_name="Parámetros usados")
        from openpyxl.styles import Font
        for ws in w.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(max(len(str(col[0].value or "")) + 2, 12), 40)
        ws = w.book["Liquidación"]
        for fila in ws.iter_rows(min_row=2):
            for cell in fila:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.##"
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    return buffer.getvalue()


def _flash():
    f = st.session_state.pop("as_flash", None)
    if f:
        getattr(st, f[0])(f[1])


def _set_flash(tipo, texto):
    st.session_state["as_flash"] = (tipo, texto)


# ----------------------------------------------------------------
# Pantalla
# ----------------------------------------------------------------
def render_asistencia(usuario):
    st.title("🕒 Asistencia y Horas Extras")
    st.write("Cargá la asistencia del personal, y la app calcula las horas ordinarias, extras, de domingo y de feriado y lo que corresponde pagar.")
    if usuario.get("rol") != "admin":
        st.warning("Esta pantalla está disponible solo para el administrador (contiene salarios).")
        return
    try:
        empleados, modalidades = db.listar_empleados(), db.listar_modalidades()
        feriados_lista, params = db.listar_feriados(), db.leer_parametros()
    except Exception:
        st.warning("Todavía no están creadas las tablas de asistencia. Corré el archivo `asistencia.sql` en Supabase "
                   "(SQL Editor → New query → pegar → Run) y volvé a entrar.")
        return
    feriados = {A.a_fecha(f["fecha"]) for f in feriados_lista if A.a_fecha(f["fecha"])}
    _flash()

    tabs = st.tabs(["📝 Cargar", "📋 Registro", "💰 Liquidación", "👥 Empleados", "⚙️ Parámetros"])
    with tabs[0]:
        _tab_cargar(empleados, modalidades, feriados, feriados_lista, usuario)
    with tabs[1]:
        _tab_registro(empleados, modalidades, feriados, params)
    with tabs[2]:
        _tab_liquidacion(empleados, modalidades, feriados, params)
    with tabs[3]:
        _tab_empleados(empleados, params)
    with tabs[4]:
        _tab_parametros(params, modalidades, feriados_lista, empleados, usuario)


# ----------------------------------------------------------------
# 1. Cargar (reemplaza las macros de la planilla)
# ----------------------------------------------------------------
def _guardar_lote(registros, reemplazar, usuario):
    """Guarda los registros; si no se pide reemplazar, omite los que ya existían. Devuelve (nuevos, reemplazados, omitidos)."""
    if not registros:
        return 0, 0, 0
    fechas = sorted(r["fecha"] for r in registros)
    ya = {(r["empleado_id"], r["fecha"]) for r in db.listar_registros(fechas[0], fechas[-1])}
    nuevos = [r for r in registros if (r["empleado_id"], r["fecha"]) not in ya]
    existentes = [r for r in registros if (r["empleado_id"], r["fecha"]) in ya]
    a_guardar = nuevos + (existentes if reemplazar else [])
    for r in a_guardar:
        r["creado_por"] = usuario.get("nombre")
    db.guardar_registros(a_guardar)
    return len(nuevos), (len(existentes) if reemplazar else 0), (0 if reemplazar else len(existentes))


def _mensaje_lote(n, r, o):
    partes = [f"{n} registro(s) nuevo(s)"]
    if r:
        partes.append(f"{r} reemplazado(s)")
    if o:
        partes.append(f"{o} omitido(s) porque ya estaban cargados (tildá «Reemplazar» si querés pisarlos)")
    return "Listo: " + ", ".join(partes) + "."


def _selector_personal(empleados, clave):
    activos = [e for e in empleados if e.get("activo") is not False]
    areas = sorted({e.get("area") or "Sin área" for e in activos}, key=str.casefold)
    elegidas = st.multiselect("Áreas", areas, default=areas, key=f"{clave}_areas")
    candidatos = [e for e in activos if (e.get("area") or "Sin área") in elegidas]
    nombres = [e["nombre"] for e in candidatos]
    marcados = st.multiselect("Personal", nombres, default=nombres, key=f"{clave}_emps_{'|'.join(elegidas)}")
    return [e for e in candidatos if e["nombre"] in marcados]


def _tab_cargar(empleados, modalidades, feriados, feriados_lista, usuario):
    mods = A.modalidad_por_nombre(modalidades)
    nombres_mod = [m["nombre"] for m in modalidades]
    motivo = {A.a_fecha(f["fecha"]): f.get("motivo") for f in feriados_lista}
    if not [e for e in empleados if e.get("activo") is not False]:
        st.info("Primero cargá a los empleados en la pestaña «Empleados» (o importá tu planilla en «Parámetros»).")
        return
    if not nombres_mod:
        st.info("No hay modalidades de turno cargadas (pestaña «Parámetros»).")
        return

    def descripcion_dia(d):
        tipo = A.tipo_de_dia(d, feriados)
        etiqueta = {"HABIL": "día hábil", "DOMINGO": "DOMINGO", "FERIADO": f"FERIADO ({motivo.get(d) or 'sin motivo'})"}[tipo]
        return f"{_fecha_larga(d)} · {etiqueta}"

    st.markdown("##### 📅 Un día de toda la cuadrilla")
    c1, c2, c3 = st.columns(3)
    fecha = c1.date_input("Fecha", value=date.today(), format="DD/MM/YYYY", key="as_dia_fecha")
    mod = c2.selectbox("Turno por defecto", nombres_mod, key="as_dia_mod")
    semana = c3.text_input("Etiqueta de semana (opcional)", placeholder="SEMANA 1", key="as_dia_sem")
    st.caption(descripcion_dia(fecha))
    personal = _selector_personal(empleados, "as_dia")
    reemplazar = st.checkbox("Reemplazar si ya hay un registro ese día", key="as_dia_reemp")
    if st.button(f"Cargar el día ({len(personal)} persona(s))", type="primary", key="as_dia_btn", disabled=not personal):
        regs = [A.registro_por_defecto(e["id"], fecha, mods[mod], semana) for e in personal]
        n, r, o = _guardar_lote(regs, reemplazar, usuario)
        _set_flash("success", _mensaje_lote(n, r, o))
        st.rerun()

    st.markdown("---")
    st.markdown("##### 🗓️ Una semana completa")
    d1, d2, d3 = st.columns(3)
    hoy = date.today()
    lunes = d1.date_input("Lunes de la semana", value=hoy - timedelta(days=hoy.weekday()), format="DD/MM/YYYY", key="as_sem_lunes")
    mod_s = d2.selectbox("Turno por defecto", nombres_mod, key="as_sem_mod")
    con_domingo = d3.checkbox("Incluir el domingo", key="as_sem_dom")
    semana_s = st.text_input("Etiqueta de semana (opcional)", placeholder="SEMANA 1", key="as_sem_et")
    if lunes.weekday() != 0:
        st.warning(f"La fecha elegida es {DIAS[lunes.weekday()]}, no lunes. Se cargará desde esa fecha.")
    dias = A.fechas_de_semana(lunes, 7 if con_domingo else 6)
    st.caption("Se cargan: " + ", ".join(f"{DIAS[d.weekday()][:3]} {d:%d/%m}" for d in dias))
    personal_s = _selector_personal(empleados, "as_sem")
    reemplazar_s = st.checkbox("Reemplazar los registros que ya existan", key="as_sem_reemp")
    if st.button(f"Cargar la semana ({len(personal_s)} persona(s) × {len(dias)} días)", type="primary", key="as_sem_btn", disabled=not personal_s):
        regs = [A.registro_por_defecto(e["id"], d, mods[mod_s], semana_s) for d in dias for e in personal_s]
        n, r, o = _guardar_lote(regs, reemplazar_s, usuario)
        _set_flash("success", _mensaje_lote(n, r, o))
        st.rerun()

    st.markdown("---")
    st.markdown("##### 🛡️ Guardia individual (domingo, feriado o turno especial)")
    g1, g2 = st.columns(2)
    fecha_g = g1.date_input("Fecha de la guardia", value=date.today(), format="DD/MM/YYYY", key="as_gi_fecha")
    mod_g = g2.selectbox("Turno", nombres_mod, key="as_gi_mod")
    st.caption(descripcion_dia(fecha_g))
    base = mods[mod_g]
    h1, h2, h3, h4, h5 = st.columns(5)
    ing = h1.time_input("Ingreso", value=A.a_hora(base.get("ingreso")) or time(6, 0), key=f"as_gi_ing_{mod_g}", step=900)
    sal = h2.time_input("Salida", value=A.a_hora(base.get("salida")) or time(18, 0), key=f"as_gi_sal_{mod_g}", step=900)
    desc = h3.number_input("Descanso (hs)", min_value=0.0, value=float(base.get("descanso") or 0), step=0.5, key=f"as_gi_desc_{mod_g}")
    adic_d = h4.number_input("HE adicional diurna", min_value=0.0, value=0.0, step=0.5, key="as_gi_ad")
    adic_n = h5.number_input("HE adicional nocturna", min_value=0.0, value=0.0, step=0.5, key="as_gi_an")
    obs = st.text_input("Observaciones", placeholder="Ej: Guardia por mantenimiento de parada", key="as_gi_obs")
    activos = [e for e in empleados if e.get("activo") is not False]
    quienes = st.multiselect("Personal que asistió", [e["nombre"] for e in activos], key="as_gi_emps")
    reemplazar_g = st.checkbox("Reemplazar si ya hay un registro ese día", value=True, key="as_gi_reemp")
    if st.button(f"Guardar la guardia ({len(quienes)} persona(s))", type="primary", key="as_gi_btn", disabled=not quienes):
        trabaja = base.get("estado") == "Trabajo"
        regs = [{"empleado_id": e["id"], "fecha": fecha_g.isoformat(), "modalidad": mod_g, "ingreso": ing.strftime("%H:%M") if trabaja else None,
                 "salida": sal.strftime("%H:%M") if trabaja else None, "descanso": desc if trabaja else 0.0, "he_diurna_adic": adic_d,
                 "he_nocturna_adic": adic_n, "observaciones": obs.strip() or "Guardia", "semana": None}
                for e in activos if e["nombre"] in quienes]
        n, r, o = _guardar_lote(regs, reemplazar_g, usuario)
        _set_flash("success", _mensaje_lote(n, r, o))
        st.rerun()


# ----------------------------------------------------------------
# 2. Registro (editable)
# ----------------------------------------------------------------
def _tab_registro(empleados, modalidades, feriados, params):
    mods = A.modalidad_por_nombre(modalidades)
    por_id = {e["id"]: e for e in empleados}
    hoy = date.today()
    f1, f2, f3, f4 = st.columns(4)
    desde = f1.date_input("Desde", value=hoy.replace(day=1), format="DD/MM/YYYY", key="as_reg_desde")
    hasta = f2.date_input("Hasta", value=hoy, format="DD/MM/YYYY", key="as_reg_hasta")
    areas = sorted({e.get("area") or "Sin área" for e in empleados}, key=str.casefold)
    f_area = f3.selectbox("Área", ["Todas"] + areas, key="as_reg_area")
    f_emp = f4.selectbox("Empleado", ["Todos"] + sorted([e["nombre"] for e in empleados], key=str.casefold), key="as_reg_emp")
    solo_avisos = st.checkbox("Mostrar solo los días con avisos (por ejemplo, más horas extra que el tope)", key="as_reg_avisos")
    if hasta < desde:
        st.error("La fecha «Hasta» no puede ser anterior a «Desde».")
        return
    registros = db.listar_registros(desde, hasta)
    filas, originales = [], {}
    for r in registros:
        e = por_id.get(r["empleado_id"])
        if not e or (f_area != "Todas" and (e.get("area") or "Sin área") != f_area) or (f_emp != "Todos" and e["nombre"] != f_emp):
            continue
        c = A.calcular_registro(r, mods, params, feriados)
        if solo_avisos and not c["avisos"]:
            continue
        d = A.a_fecha(r["fecha"])
        originales[r["id"]] = {k: r.get(k) for k in CAMPOS_REGISTRO}
        dia = TIPO_DIA_TXT[c["tipo_dia"]].capitalize() + (" → " + TIPO_DIA_TXT[c["segmentos"][-1]["tipo_dia"]] if c["cruza"] and c["segmentos"][-1]["tipo_dia"] != c["tipo_dia"] else "")
        filas.append({"id": r["id"], "Fecha": f"{d:%d/%m/%Y} {DIAS[d.weekday()][:3]}", "Empleado": e["nombre"], "Modalidad": r.get("modalidad"),
                      "Ingreso": A.a_hora(r.get("ingreso")), "Salida": A.a_hora(r.get("salida")), "Descanso": float(r.get("descanso") or 0),
                      "HE adic. diurna": float(r.get("he_diurna_adic") or 0), "HE adic. nocturna": float(r.get("he_nocturna_adic") or 0),
                      "Observaciones": r.get("observaciones") or "", "Día": dia, "Netas": c["netas"], "Ord.": c["ordinarias"],
                      "HE diurnas": c["horas"]["he_diurna"], "HE noct.": c["horas"]["he_nocturna"],
                      "Domingo": c["horas"]["domingo_diurno"] + c["horas"]["domingo_nocturno"],
                      "Feriado": c["horas"]["feriado_diurno"] + c["horas"]["feriado_nocturno"], "Avisos": " ".join(c["avisos"])})
    if not filas:
        st.info("No hay registros con esos filtros en el período.")
        return
    LIMITE = 600
    if len(filas) > LIMITE:
        st.warning(f"Hay {len(filas)} registros: se muestran los primeros {LIMITE}. Acotá el período o filtrá por empleado.")
        filas = filas[:LIMITE]
        originales = {f["id"]: originales[f["id"]] for f in filas}
    st.caption(f"{len(filas)} registro(s). Las columnas con ✏️ se editan (doble click); las demás se calculan solas al guardar.")
    df = pd.DataFrame(filas)
    firma = f"{desde}|{hasta}|{f_area}|{f_emp}|{solo_avisos}|{st.session_state.get('as_reg_version', 0)}"
    editables = ["Modalidad", "Ingreso", "Salida", "Descanso", "HE adic. diurna", "HE adic. nocturna", "Observaciones"]
    editado = st.data_editor(
        df, key=f"as_reg_editor_{firma}", hide_index=True, width="stretch", height=min(70 + 36 * len(df), 560),
        disabled=[c for c in df.columns if c not in editables],
        column_config={"id": None,
                       "Modalidad": st.column_config.SelectboxColumn("Modalidad ✏️", options=list(mods.keys())),
                       "Ingreso": st.column_config.TimeColumn("Ingreso ✏️", format="HH:mm", step=300),
                       "Salida": st.column_config.TimeColumn("Salida ✏️", format="HH:mm", step=300),
                       "Descanso": st.column_config.NumberColumn("Descanso ✏️", min_value=0.0, step=0.5, format="%.1f"),
                       "HE adic. diurna": st.column_config.NumberColumn("HE adic. diurna ✏️", min_value=0.0, step=0.5, format="%.1f"),
                       "HE adic. nocturna": st.column_config.NumberColumn("HE adic. noct. ✏️", min_value=0.0, step=0.5, format="%.1f"),
                       "Observaciones": st.column_config.TextColumn("Observaciones ✏️")})
    nuevos = {int(r["id"]): {k: r[c] for k, c in zip(CAMPOS_REGISTRO, ["Modalidad", "Ingreso", "Salida", "Descanso", "HE adic. diurna", "HE adic. nocturna", "Observaciones"])}
              for _, r in editado.iterrows()}
    cambios = calcular_cambios_registros(originales, nuevos)
    if st.button(f"💾 Guardar {len(cambios)} cambio(s)", type="primary", disabled=not cambios, key="as_reg_guardar"):
        errores = 0
        for rid, parche in cambios:
            try:
                db.actualizar_registro(rid, parche)
            except Exception:
                errores += 1
        st.session_state["as_reg_version"] = st.session_state.get("as_reg_version", 0) + 1
        _set_flash("success" if not errores else "warning", f"Se guardaron {len(cambios) - errores} cambio(s)." + (f" {errores} fallaron." if errores else ""))
        st.rerun()

    with st.expander("🗑️ Quitar registros"):
        etiquetas = {f"{f['Fecha']} · {f['Empleado']}": f["id"] for f in filas}
        elegidos = st.multiselect("Registros a quitar", list(etiquetas.keys()), key="as_reg_quitar_sel")
        if elegidos and st.checkbox(f"Confirmo que quiero quitar {len(elegidos)} registro(s)", key="as_reg_quitar_ok") and st.button("Quitar", key="as_reg_quitar_btn"):
            db.borrar_registros([etiquetas[e] for e in elegidos])
            st.session_state["as_reg_version"] = st.session_state.get("as_reg_version", 0) + 1
            _set_flash("success", f"Se quitaron {len(elegidos)} registro(s).")
            st.rerun()


# ----------------------------------------------------------------
# 3. Liquidación
# ----------------------------------------------------------------
def _tab_liquidacion(empleados, modalidades, feriados, params):
    hoy = date.today()
    c1, c2, c3 = st.columns(3)
    desde = c1.date_input("Desde (fecha de corte)", value=hoy.replace(day=1), format="DD/MM/YYYY", key="as_liq_desde")
    hasta = c2.date_input("Hasta (fecha de corte)", value=hoy, format="DD/MM/YYYY", key="as_liq_hasta")
    areas = sorted({e.get("area") or "Sin área" for e in empleados}, key=str.casefold)
    f_area = c3.selectbox("Área", ["Todas"] + areas, key="as_liq_area")
    if hasta < desde:
        st.error("La fecha «Hasta» no puede ser anterior a «Desde».")
        return
    elegidos = [e for e in empleados if f_area == "Todas" or (e.get("area") or "Sin área") == f_area]
    registros = db.listar_registros(desde, hasta)
    liq = A.liquidar(elegidos, registros, modalidades, params, feriados, desde, hasta)
    if not liq["filas"]:
        st.info("No hay registros en ese período.")
        return
    t = liq["totales"]
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Total a pagar por horas extras", _fmt_gs(t["total"]))
    m2.metric("HE diurnas + nocturnas", _hs(round(t["he_diurna_horas"] + t["he_nocturna_horas"], 2)) + " hs")
    m3.metric("Horas de domingo", _hs(round(t["domingo_diurno_horas"] + t["domingo_nocturno_horas"], 2)) + " hs")
    m4.metric("Horas de feriado", _hs(round(t["feriado_diurno_horas"] + t["feriado_nocturno_horas"], 2)) + " hs")

    def fila_tabla(f):
        return {"Empleado": f["empleado"], "Área": f["area"], "Salario": f["salario"], "Días": f["dias_trabajados"], "Aus.": f["ausencias"],
                "HE diurnas (hs)": f["he_diurna_horas"], "HE diurnas (Gs.)": f["he_diurna_monto"],
                "HE noct. (hs)": f["he_nocturna_horas"], "HE noct. (Gs.)": f["he_nocturna_monto"],
                "Domingo (hs)": round(f["domingo_diurno_horas"] + f["domingo_nocturno_horas"], 2), "Domingo (Gs.)": f["domingo_diurno_monto"] + f["domingo_nocturno_monto"],
                "Feriado (hs)": round(f["feriado_diurno_horas"] + f["feriado_nocturno_horas"], 2), "Feriado (Gs.)": f["feriado_diurno_monto"] + f["feriado_nocturno_monto"],
                "TOTAL (Gs.)": f["total"]}
    tabla = [fila_tabla(f) for f in liq["filas"]]
    tabla.append(fila_tabla({**t, "empleado": "TOTAL", "area": ""}))
    st.dataframe(pd.DataFrame(tabla), hide_index=True, width="stretch",
                 column_config={c: st.column_config.NumberColumn(format="%d") for c in ["Salario", "HE diurnas (Gs.)", "HE noct. (Gs.)", "Domingo (Gs.)", "Feriado (Gs.)", "TOTAL (Gs.)"]})

    if liq["avisos"]:
        st.markdown("##### Avisos")
        for a in liq["avisos"][:15]:
            st.markdown(f"{ICONO[a['nivel']]} {a['texto']}")
        if len(liq["avisos"]) > 15:
            with st.expander(f"Ver los {len(liq['avisos']) - 15} avisos restantes"):
                for a in liq["avisos"][15:]:
                    st.markdown(f"{ICONO[a['nivel']]} {a['texto']}")

    with st.expander("🔎 Detalle día por día de un empleado"):
        nombres = {f["empleado"]: f["empleado_id"] for f in liq["filas"]}
        quien = st.selectbox("Empleado", list(nombres.keys()), key="as_liq_det")
        filas = []
        for d in liq["detalle"].get(nombres[quien], []):
            c = d["calculo"]
            filas.append({"Fecha": f"{d['fecha']:%d/%m/%Y} {DIAS[d['fecha'].weekday()][:3]}", "Modalidad": d["modalidad"],
                          "Tramos": _tramos_txt(c),
                          "Netas": c["netas"], "Ord.": c["ordinarias"], "HE d.": c["horas"]["he_diurna"], "HE n.": c["horas"]["he_nocturna"],
                          "Dom. d.": c["horas"]["domingo_diurno"], "Dom. n.": c["horas"]["domingo_nocturno"],
                          "Fer. d.": c["horas"]["feriado_diurno"], "Fer. n.": c["horas"]["feriado_nocturno"], "Avisos": " ".join(c["avisos"])})
        st.dataframe(pd.DataFrame(filas), hide_index=True, width="stretch")

    excel = construir_excel(liq, {e["id"]: e for e in empleados}, params, desde, hasta)
    st.download_button("⬇️ Descargar la liquidación (Excel con detalle día por día)", data=excel, key="as_liq_excel",
                       file_name=f"liquidacion_horas_extras_{desde:%Y%m%d}_{hasta:%Y%m%d}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.caption("Guardá este Excel al cerrar cada período: si después cambiás salarios o parámetros, la pantalla recalcula con los valores nuevos.")


# ----------------------------------------------------------------
# 4. Empleados
# ----------------------------------------------------------------
def _tab_empleados(empleados, params):
    st.caption("🔒 Esta lista contiene salarios: solo el administrador la ve. Incluí al personal de mantenimiento y de las demás áreas.")
    areas_existentes = sorted({e.get("area") for e in empleados if e.get("area")}, key=str.casefold)
    with st.expander("➕ Agregar empleado"):
        with st.form("as_form_emp", clear_on_submit=True):
            c1, c2 = st.columns(2)
            nombre = c1.text_input("Nombre y apellido *")
            area = c2.text_input("Área", placeholder="Mantenimiento, Producción, Calidad…" + (f" (ya usadas: {', '.join(areas_existentes)})" if areas_existentes else ""))
            c3, c4 = st.columns(2)
            salario = c3.number_input("Salario base (Gs.) *", min_value=0, step=10000, value=0)
            adicional = c4.number_input("Adicional fijo (Gs.)", min_value=0, step=10000, value=0)
            if st.form_submit_button("Guardar empleado"):
                if not nombre.strip():
                    st.error("El nombre es obligatorio.")
                elif any(str(e["nombre"]).strip().casefold() == nombre.strip().casefold() for e in empleados):
                    st.error("Ya existe un empleado con ese nombre.")
                else:
                    db.insertar_empleado({"nombre": nombre.strip(), "area": area.strip() or None, "salario_base": salario, "adicional_fijo": adicional, "activo": True})
                    _set_flash("success", f"{nombre.strip()} agregado.")
                    st.rerun()
    if not empleados:
        st.info("Todavía no hay empleados. Agregalos acá o importá tu planilla en «Parámetros».")
        return
    sin_salario = [e["nombre"] for e in empleados if e.get("activo") is not False and float(e.get("salario_base") or 0) + float(e.get("adicional_fijo") or 0) <= 0]
    if sin_salario:
        st.warning("Sin salario cargado (sus horas extra no se pueden valorizar): " + ", ".join(sin_salario))
    df = pd.DataFrame([{"id": e["id"], "Nombre": e["nombre"], "Área": e.get("area") or "", "Salario base (Gs.)": float(e.get("salario_base") or 0),
                        "Adicional fijo (Gs.)": float(e.get("adicional_fijo") or 0), "Valor hora (Gs.)": round(A.valor_hora(e, params), 2),
                        "Activo": e.get("activo") is not False} for e in empleados])
    firma = f"{len(empleados)}|{st.session_state.get('as_emp_version', 0)}"
    editado = st.data_editor(df, key=f"as_emp_editor_{firma}", hide_index=True, width="stretch", disabled=["id", "Valor hora (Gs.)"],
                             column_config={"id": None, "Salario base (Gs.)": st.column_config.NumberColumn(min_value=0, step=10000, format="%d"),
                                            "Adicional fijo (Gs.)": st.column_config.NumberColumn(min_value=0, step=10000, format="%d"),
                                            "Valor hora (Gs.)": st.column_config.NumberColumn(format="%.2f"), "Activo": st.column_config.CheckboxColumn()})
    cambios, error = [], None
    nombres_vistos = set()
    por_id = {e["id"]: e for e in empleados}
    for _, r in editado.iterrows():
        nombre = str(r["Nombre"]).strip()
        if not nombre or nombre.casefold() in nombres_vistos:
            error = "Hay nombres vacíos o repetidos: corregilos para poder guardar."
        nombres_vistos.add(nombre.casefold())
        e = por_id[int(r["id"])]
        nuevo = {"nombre": nombre, "area": str(r["Área"]).strip() or None, "salario_base": float(r["Salario base (Gs.)"]), "adicional_fijo": float(r["Adicional fijo (Gs.)"]),
                 "activo": bool(r["Activo"])}
        previo = {"nombre": e["nombre"], "area": e.get("area") or None, "salario_base": float(e.get("salario_base") or 0), "adicional_fijo": float(e.get("adicional_fijo") or 0),
                  "activo": e.get("activo") is not False}
        parche = {k: v for k, v in nuevo.items() if v != previo[k]}
        if parche:
            cambios.append((int(r["id"]), parche))
    if error:
        st.error(error)
    if st.button(f"💾 Guardar {len(cambios)} cambio(s)", type="primary", disabled=not cambios or bool(error), key="as_emp_guardar"):
        errores = 0
        for eid, parche in cambios:
            try:
                db.actualizar_empleado(eid, parche)
            except Exception:
                errores += 1
        st.session_state["as_emp_version"] = st.session_state.get("as_emp_version", 0) + 1
        _set_flash("success" if not errores else "warning", f"Se guardaron {len(cambios) - errores} cambio(s)." + (f" {errores} fallaron." if errores else ""))
        st.rerun()
    st.caption("Los empleados con registros no se pueden borrar (se perdería el historial): destildá «Activo» para dejar de usarlos.")


# ----------------------------------------------------------------
# 5. Parámetros, modalidades, feriados e importación
# ----------------------------------------------------------------
def _tab_parametros(params, modalidades, feriados_lista, empleados, usuario):
    st.markdown("##### Reglas de cálculo")
    with st.form("as_form_params"):
        claves = list(NOMBRE_PARAM.keys())
        cols = st.columns(3)
        nuevos = {}
        for i, k in enumerate(claves):
            nuevos[k] = cols[i % 3].number_input(NOMBRE_PARAM[k], min_value=0.0, value=float(params[k]), step=0.1 if k.startswith("factor") else 1.0, key=f"as_par_{k}")
        domingo = st.checkbox("Todo el turno de domingo o feriado se paga con su factor (sin contar horas ordinarias)", value=bool(params["domingo_todo_el_turno"]), key="as_par_dom")
        cruce = st.checkbox("Turnos que cruzan la medianoche: pagar por hora real (cada hora según el día en que cae). "
                            "Si lo destildás, todo el turno se trata según el día en que empieza, como la planilla anterior.", value=bool(params["cruce_por_hora_real"]), key="as_par_cruce")
        if st.form_submit_button("Guardar reglas", type="primary"):
            db.guardar_parametros({**nuevos, "domingo_todo_el_turno": 1.0 if domingo else 0.0, "cruce_por_hora_real": 1.0 if cruce else 0.0})
            _set_flash("success", "Reglas guardadas. La liquidación se recalcula con los valores nuevos.")
            st.rerun()
    st.caption("El valor hora es: (salario base + adicional fijo) ÷ días base ÷ horas de la jornada diurna. Los factores multiplican ese valor: "
               "1,5 = +50%, 2,0 = +100%, 2,3 = +130%.")

    st.markdown("---")
    st.markdown("##### Modalidades de turno")
    df = pd.DataFrame([{"Nombre": m["nombre"], "Ingreso": A.a_hora(m.get("ingreso")), "Salida": A.a_hora(m.get("salida")), "Descanso (hs)": float(m.get("descanso") or 0),
                        "Tipo": m.get("tipo"), "Estado": m.get("estado")} for m in modalidades])
    editado = st.data_editor(df, key=f"as_mod_editor_{len(modalidades)}_{st.session_state.get('as_mod_version', 0)}", hide_index=True, width="stretch", disabled=["Nombre"],
                             column_config={"Ingreso": st.column_config.TimeColumn(format="HH:mm", step=300), "Salida": st.column_config.TimeColumn(format="HH:mm", step=300),
                                            "Tipo": st.column_config.SelectboxColumn(options=A.TIPOS_RECARGO), "Estado": st.column_config.SelectboxColumn(options=A.ESTADOS_MODALIDAD)})
    cambios = []
    originales = {m["nombre"]: m for m in modalidades}
    for _, r in editado.iterrows():
        m = originales[r["Nombre"]]
        nuevo = {"nombre": r["Nombre"], "ingreso": (_hora_celda(r["Ingreso"]).strftime("%H:%M") if _hora_celda(r["Ingreso"]) else None),
                 "salida": (_hora_celda(r["Salida"]).strftime("%H:%M") if _hora_celda(r["Salida"]) else None), "descanso": float(r["Descanso (hs)"]),
                 "tipo": r["Tipo"], "estado": r["Estado"]}
        previo = {"nombre": m["nombre"], "ingreso": A.hora_texto(m.get("ingreso")) or None, "salida": A.hora_texto(m.get("salida")) or None,
                  "descanso": float(m.get("descanso") or 0), "tipo": m.get("tipo"), "estado": m.get("estado")}
        if nuevo != previo:
            cambios.append(nuevo)
    if st.button(f"💾 Guardar {len(cambios)} modalidad(es)", disabled=not cambios, key="as_mod_guardar"):
        for m in cambios:
            db.guardar_modalidad(m)
        st.session_state["as_mod_version"] = st.session_state.get("as_mod_version", 0) + 1
        _set_flash("success", f"Se guardaron {len(cambios)} modalidad(es).")
        st.rerun()
    with st.expander("➕ Agregar una modalidad"):
        with st.form("as_form_mod", clear_on_submit=True):
            c1, c2, c3 = st.columns(3)
            nombre = c1.text_input("Nombre *", placeholder="Ej: Tarde")
            tipo = c2.selectbox("Tipo de recargo", A.TIPOS_RECARGO)
            estado = c3.selectbox("Estado", A.ESTADOS_MODALIDAD)
            d1, d2, d3 = st.columns(3)
            ing = d1.time_input("Ingreso", value=time(14, 0), step=900)
            sal = d2.time_input("Salida", value=time(22, 0), step=900)
            desc = d3.number_input("Descanso (hs)", min_value=0.0, value=1.0, step=0.5)
            if st.form_submit_button("Guardar modalidad"):
                if not nombre.strip():
                    st.error("El nombre es obligatorio.")
                elif nombre.strip().casefold() in {m["nombre"].casefold() for m in modalidades}:
                    st.error("Ya existe una modalidad con ese nombre.")
                else:
                    trabaja = estado == "Trabajo"
                    db.guardar_modalidad({"nombre": nombre.strip(), "ingreso": ing.strftime("%H:%M") if trabaja else None, "salida": sal.strftime("%H:%M") if trabaja else None,
                                          "descanso": desc if trabaja else 0.0, "tipo": tipo if trabaja else "NINGUNO", "estado": estado})
                    _set_flash("success", "Modalidad agregada.")
                    st.rerun()

    st.markdown("---")
    st.markdown("##### Feriados")
    if feriados_lista:
        st.dataframe(pd.DataFrame([{"Fecha": f"{A.a_fecha(f['fecha']):%d/%m/%Y} {DIAS[A.a_fecha(f['fecha']).weekday()][:3]}", "Motivo": f.get("motivo") or ""} for f in feriados_lista]),
                     hide_index=True, width="stretch", height=min(70 + 36 * len(feriados_lista), 320))
    else:
        st.info("No hay feriados cargados.")
    st.caption("⚠️ Revisá esta lista contra el calendario oficial de cada año: un feriado que falte se calcula como día común.")
    f1, f2, f3 = st.columns([0.25, 0.5, 0.25])
    nueva_fecha = f1.date_input("Fecha del feriado", value=None, format="DD/MM/YYYY", key="as_fer_fecha")
    motivo = f2.text_input("Motivo", key="as_fer_motivo")
    if f3.button("Agregar feriado", key="as_fer_btn", disabled=nueva_fecha is None):
        db.guardar_feriado(nueva_fecha, motivo.strip())
        _set_flash("success", f"Feriado del {nueva_fecha:%d/%m/%Y} guardado.")
        st.rerun()
    if feriados_lista:
        opciones = {f"{A.a_fecha(f['fecha']):%d/%m/%Y} · {f.get('motivo') or ''}": f["fecha"] for f in feriados_lista}
        quitar = st.selectbox("Quitar un feriado", ["—"] + list(opciones.keys()), key="as_fer_quitar")
        if quitar != "—" and st.button("Quitar feriado", key="as_fer_quitar_btn"):
            db.borrar_feriado(opciones[quitar])
            _set_flash("success", "Feriado quitado.")
            st.rerun()

    st.markdown("---")
    st.markdown("##### 📥 Importar la planilla de Excel anterior")
    st.caption("Sube tu archivo «Sistema de Control de Asistencia y Horas Extras» (.xlsm o .xlsx). Se leen los empleados, las modalidades, los feriados, los parámetros y todos los "
               "registros diarios. No se pisan los empleados ni los salarios que ya estén cargados.")
    archivo = st.file_uploader("Planilla de Excel", type=["xlsm", "xlsx"], key="as_import_file")
    if archivo is not None:
        datos = A.leer_planilla(archivo.getvalue())
        if datos.get("error"):
            st.error(datos["error"])
            return
        fechas = sorted(r["fecha"] for r in datos["registros"])
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Empleados", len(datos["empleados"]))
        c2.metric("Registros diarios", len(datos["registros"]))
        c3.metric("Feriados", len(datos["feriados"]))
        c4.metric("Período", f"{A.a_fecha(fechas[0]):%d/%m} al {A.a_fecha(fechas[-1]):%d/%m/%Y}" if fechas else "—")
        for aviso in datos["avisos"]:
            st.warning(aviso)
        area = st.text_input("Área para los empleados nuevos", value="Mantenimiento", key="as_import_area",
                             help="Después podés cambiarla empleado por empleado en la pestaña «Empleados».")
        reemplazar = st.checkbox("Reemplazar los registros de esos días que ya existan en la app", value=True, key="as_import_reemp")
        con_params = st.checkbox("Traer también los parámetros de la planilla (jornada y factores de hora extra)", value=True, key="as_import_params",
                                 help="Si ya ajustaste las reglas en la app y no querés perderlas, destildalo.")
        st.info("Por defecto la app paga el domingo nocturno y el feriado nocturno al 130% (factor 2,3) y reparte los turnos que cruzan la medianoche por hora real "
                "(la planilla no distinguía eso); podés cambiarlo en «Reglas de cálculo».")
        if st.button("📥 Importar a la app", type="primary", key="as_import_btn"):
            r = db.importar_planilla(datos, area.strip() or None, reemplazar, usuario.get("nombre"), con_params)
            _set_flash("success", f"Importado: {r['empleados_nuevos']} empleado(s) nuevo(s) ({r['empleados_existentes']} ya existían), "
                                  f"{r['registros_nuevos']} registro(s) nuevo(s), {r['registros_reemplazados']} reemplazado(s), {r['registros_omitidos']} omitido(s), "
                                  f"{r['feriados']} feriado(s) y {r['modalidades']} modalidad(es).")
            st.rerun()
